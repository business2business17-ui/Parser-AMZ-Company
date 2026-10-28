"""Amazon seller research assistant. Research only on pages you are permitted to access.
Run: python agent.py --input seeds.csv --out reports
Or: python agent.py --asin B0D9VS4PV7 --marketplace US --out reports
No CAPTCHAs, authentication, geolocation or bot defenses are bypassed.
"""
from __future__ import annotations
import argparse, csv, json, re, time, logging, os, sys, tempfile
import math
from dataclasses import dataclass, asdict, field
from datetime import datetime, timezone
from pathlib import Path
from urllib.parse import urlparse, parse_qs, quote, urljoin
from bs4 import BeautifulSoup
from playwright.sync_api import sync_playwright
from openpyxl import Workbook, load_workbook
from openpyxl.styles import PatternFill, Font, Alignment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.utils import get_column_letter

MARKETS = {
'US':('www.amazon.com','ATVPDKIKX0DER'), 'CA':('www.amazon.ca','A2EUQ1WTGCTBG2'),
'MX':('www.amazon.com.mx','A1AM78C64UM0Y8'), 'UK':('www.amazon.co.uk','A1F83G8C2ARO7P'),
'DE':('www.amazon.de','A1PA6795UKMFR9'), 'FR':('www.amazon.fr','A13V1IB3VIYZZH'),
'ES':('www.amazon.es','A1RKKUPIHCS9HS'), 'IT':('www.amazon.it','APJ6JRA9NG5V4'),
'PL':('www.amazon.pl','A1C3SOZRARQ6R3'), 'NL':('www.amazon.nl','A1805IZSGTT6HS'),
'BE':('www.amazon.com.be','AMEN7PMS3EDWL'), 'SE':('www.amazon.se','A2NODRKZP88ZB9'),
'AE':('www.amazon.ae','A2VIGQ35RCS4UG'), 'SA':('www.amazon.sa','A17E79C6D8DWNP'),
'IN':('www.amazon.in','A21TJRUUN4KGV'), 'AU':('www.amazon.com.au','A39IBJ37TRP1C6'),
'JP':('www.amazon.co.jp','A1VC38T7YXB528')
}
CURRENCY = {'US':'USD','CA':'CAD','MX':'MXN','UK':'GBP','DE':'EUR','FR':'EUR','ES':'EUR','IT':'EUR','PL':'PLN',
'NL':'EUR','BE':'EUR','SE':'SEK','AE':'AED','SA':'SAR','IN':'INR','AU':'AUD','JP':'JPY'}
# Default delivery postal codes (central city of each country). AE/SA choose a city instead of a code.
ZIPS = {'US':'10001','CA':'M5V 2T6','MX':'06600','UK':'SW1A 1AA','DE':'10115','FR':'75001','ES':'28001',
'IT':'00100','PL':'00-001','NL':'1012 AB','BE':'1000','SE':'111 20','IN':'110001','AU':'2000','JP':'100-0001'}
CITY_MARKETS = {'AE':'Dubai','SA':'Riyadh'}
ASIN_RE = re.compile(r'(?<![A-Z0-9])([A-Z0-9]{10})(?![A-Z0-9])', re.I)
DEVICE_RE = re.compile(r'(?i)\b(?:iPhone\s*(?:1[1-9]|[2-9][0-9]?)(?:\s*(?:Pro\s*Max|Pro|Plus|Mini))?|AirPods\s*(?:Pro\s*[123]?|[1234](?:st|nd|rd|th)?\s*Gen(?:eration)?|Max))\b')
PRICE_RE = re.compile(r'(?<!\d)(\d{1,4}(?:[.,]\d{2})?)(?!\d)')
ISO_RE = re.compile(r'\b(?:US|CA|MX|GB|UK|DE|FR|ES|IT|PL|NL|BE|SE|AE|SA|IN|AU|JP|CN|HK|SG|TR|VN|KR|TW|BR)\b')
EMAIL_RE = re.compile(r'\b[A-Z0-9._%+\-]+@[A-Z0-9.\-]+\.[A-Z]{2,}\b', re.I)
PHONE_RE = re.compile(r'(?<!\w)(?:\+?\d[\d\s().\-]{6,}\d)(?!\w)')
REGISTRATION_RE = re.compile(r'(?i)\b(?:registration(?:\s+number|\s+no\.?)?|company\s+number|handelsregisternummer|registre\s+du\s+commerce|registro\s+mercantil|ust-?id|vat\s*(?:id|number)|eori)\s*[:#-]?\s*([A-Z0-9][A-Z0-9 ./-]{3,40})')
LOG = logging.getLogger('amazon-agent')
BASE_DIR = Path(__file__).resolve().parent

def clean(s): return re.sub(r'\s+', ' ', s or '').strip()
def txt(node): return clean(node.get_text(' ', strip=True)) if node else ''
def pick(soup, *selectors):
    for selector in selectors:
        elem = soup.select_one(selector)
        if elem and txt(elem): return txt(elem)
    return ''
def num(s):
    s = clean(s).replace('\u202f','').replace(' ','')
    m = re.search(r'\d[\d.,]*', s)
    if not m:return None
    v=m.group(0)
    if ',' in v and '.' in v: v=v.replace(',','') if v.rfind('.')>v.rfind(',') else v.replace('.','').replace(',','.')
    elif v.count(',')==1: v=v.replace(',','.') if len(v.split(',')[1])<=2 else v.replace(',','')
    elif v.count('.')>1: v=v.replace('.','')
    try: return float(v)
    except ValueError:return None

def extract_contact_data(text):
    emails=list(dict.fromkeys(EMAIL_RE.findall(text or '')))
    phones=[]
    for value in PHONE_RE.findall(text or ''):
        value=clean(value)
        digits=re.sub(r'\D','',value)
        if 7 <= len(digits) <= 16: phones.append(value)
    registrations=[]
    for match in REGISTRATION_RE.finditer(text or ''):
        value=clean(match.group(1)).rstrip('.,;')
        if value and value not in registrations:registrations.append(value)
    return emails,phones,registrations

def classify_shipping(text, seller_name=''):
    value=clean(text)
    low=value.lower()
    if not value:return 'UNKNOWN'
    if low in ('amazon','amazon.com','amazon.co.uk','amazon.de') or any(token in low for token in ('ships from amazon','versand durch amazon','expédié par amazon','spedito da amazon','enviado por amazon')):
        return 'AMAZON'
    if seller_name and low == clean(seller_name).lower():return 'SELLER'
    if any(token in low for token in ('ships from seller','versand durch den verkäufer','versand durch verkäufer','expédié par le vendeur','spedito dal venditore')):
        return 'SELLER'
    return 'SHOWN_TEXT_ONLY'

def asin_from(text):
    p=urlparse(text)
    match=re.search(r'/(?:dp|gp/product|product)/([A-Z0-9]{10})',p.path,re.I)
    if match:return match.group(1).upper()
    found=ASIN_RE.search(text.upper())
    return found.group(1).upper() if found and not text.startswith('http') else None

def market_from_url(url,default='US'):
    host=(urlparse(url).hostname or '').lower()
    for market,(domain,_) in MARKETS.items():
        if host in (domain,domain.removeprefix('www.')):return market
    return default

def url_asin(asin,market): return f'https://{MARKETS[market][0]}/dp/{asin}'
def url_seller(seller,market):return f'https://{MARKETS[market][0]}/sp?ie=UTF8&seller={quote(seller)}'
def url_store(seller,market):return f'https://{MARKETS[market][0]}/s?me={quote(seller)}&marketplaceID={MARKETS[market][1]}'

def get_specs(soup):
    values={}
    for tr in soup.select('#productDetails_detailBullets_sections1 tr, #productDetails_techSpec_section_1 tr, .prodDetTable tr'):
        k=txt(tr.select_one('th'));v=txt(tr.select_one('td'))
        if k and v:values[k]=v
    for li in soup.select('#detailBullets_feature_div li, #detailBulletsWrapper_feature_div li'):
        k=li.select_one('.a-text-bold')
        if k:
            key=txt(k).strip(' :\u200e');value=txt(li)
            if value.lower().startswith(key.lower()):value=value[len(key):].strip(' :\u200e')
            if value:values.setdefault(key,value)
    return values

def parse_product(html, asin, market):
    s=BeautifulSoup(html,'lxml')
    title=pick(s,'#productTitle','h1#title','#title')
    bullets=[txt(e) for e in s.select('#feature-bullets .a-list-item') if txt(e)]
    specs=get_specs(s)
    image=s.select_one('#landingImage, #imgBlkFront, #main-image')
    image_url=''
    if image:
        image_url=image.get('data-old-hires') or image.get('src') or ''
        if image_url.startswith('data:'): image_url=''
    if not image_url:
        meta=s.select_one('meta[property="og:image"]')
        image_url=meta.get('content','') if meta else ''
    price_text=pick(s,'#corePrice_feature_div .a-offscreen','#corePriceDisplay_desktop_feature_div .a-offscreen','#priceblock_ourprice','#priceblock_dealprice','#apex_desktop .a-price .a-offscreen')
    seller_href=''; seller_name=pick(s,'#sellerProfileTriggerId','#merchantInfoFeature_feature_div .offer-display-feature-text-message','#merchant-info a[href*="seller="]')
    for a in s.select('#sellerProfileTriggerId[href], #merchantInfoFeature_feature_div a[href*="seller="], #merchant-info a[href*="seller="]'):
        href=a.get('href','')
        if '/sp?' in href or 'seller=' in href:
            seller_href=href
            if not seller_name: seller_name=txt(a)
            break
    sm=re.search(r'[?&]seller=([A-Z0-9]+)',seller_href,re.I)
    seller_id=sm.group(1) if sm else ''
    if not seller_id:
        merchant=s.select_one('input#merchantID, input#sellerId')
        if merchant and re.fullmatch(r'[A-Z0-9]{8,20}',merchant.get('value','')):seller_id=merchant['value']
    shipping=pick(s,'#fulfillerInfoFeature_feature_div .offer-display-feature-text-message','#fulfillerInfoFeature_feature_div','#tabular-buybox-text .tabular-buybox-text')
    if not shipping:
        buybox=s.select('#tabular-buybox .tabular-buybox-text-container')
        shipping=txt(buybox[0]) if buybox else ''
    avail=pick(s,'#availability span','#availability')
    avail_l=avail.lower()
    if not title: status='UNVERIFIED'
    elif any(x in avail_l for x in ('currently unavailable','out of stock','nicht verfügbar','indisponible','agotado','non disponibile')):status='NO'
    elif any(x in avail_l for x in ('in stock','available','auf lager','en stock','disponibile','disponível')):status='YES'
    else:status='UNKNOWN'
    rating_text=pick(s,'#acrPopover .a-icon-alt','#averageCustomerReviews .a-icon-alt','#acrPopover')
    rating=num(rating_text)
    review_text=pick(s,'#acrCustomerReviewText','#acrCustomerReviewLink')
    bsr=' | '.join(v for k,v in specs.items() if 'best sellers rank' in k.lower() or 'bestseller' in k.lower() or 'rang' in k.lower())
    first_available=' | '.join(v for k,v in specs.items() if 'date first available' in k.lower() or 'first available' in k.lower() or 'mise en ligne' in k.lower())
    var_nodes=s.select('#twister .swatchAvailable, #twister .swatchSelect, #variation_color_name li, #variation_size_name li')
    variation='YES' if len(var_nodes)>1 else 'UNKNOWN' if not s.select_one('#twister') else 'NO'
    offers=pick(s,'#olp-upd-new-used a','#olp_feature_div a','#moreBuyingChoices_feature_div')
    offer_count=None
    m=re.search(r'(\d[\d.,]*)\s*(?:offers?|options?|new)',offers,re.I)
    if m:offer_count=int(num(m.group(1)) or 0)
    device= ', '.join(dict.fromkeys(m.group(0) for m in DEVICE_RE.finditer(title+' '+' '.join(bullets))))
    return {'marketplace':market,'asin':asin,'product_url':url_asin(asin,market),'image_url':image_url,
        'title':title,'highlights':' | '.join(bullets),'device_model':device,'price_text':price_text,
        'price_numeric':num(price_text),'currency':CURRENCY.get(market,''),'delivery_zip':'','buybox':'YES' if price_text and (seller_name or shipping) else 'UNKNOWN',
        'in_stock':status,'ships_from':shipping,'ships_from_entity':classify_shipping(shipping,seller_name),
        'sold_by':seller_name,'seller_id':seller_id,
        'variations':variation,'offer_count':offer_count,'rating':rating,'rating_count':int(num(review_text) or 0) if review_text else None,
        'bsr':bsr,'date_first_available':first_available,'item_details':specs,'source':'amazon_page',
        'observed_at':datetime.now(timezone.utc).isoformat(timespec='seconds'),'status':'OK' if title else 'UNVERIFIED'}

def parse_seller(html, seller_id, market):
    from seller_profile import parse_seller_profile
    return parse_seller_profile(html,seller_id,market,url_seller(seller_id,market),url_store(seller_id,market))

def collect_feedback_windows(page, seller):
    # Period-specific blocks are already in the HTML, even when a tab is hidden.
    # Never read the first rating from the whole page into every time period.
    return seller

def detect_block(page):
    t=page.content().lower()
    return any(x in t for x in ['enter the characters you see below','robot check','captcha'])

class BrowserClosedError(RuntimeError):
    """The run must stop and checkpoint instead of failing every remaining URL."""

def browser_closed(exc):
    return type(exc).__name__ == 'TargetClosedError' or 'target page, context or browser has been closed' in str(exc).lower()

def navigate(page,url,wait_ms=1200):
    try:
        response=page.goto(url,wait_until='domcontentloaded',timeout=35000)
        page.wait_for_timeout(wait_ms)
        if detect_block(page):return '', 'BLOCKED_CHALLENGE'
        return page.content(), f'HTTP_{response.status if response else "UNKNOWN"}'
    except Exception as exc:
        if browser_closed(exc):
            raise BrowserClosedError('Browser closed or crashed. Progress saved; resume this run with --resume.') from exc
        return '',f'ERROR:{type(exc).__name__}:{str(exc)[:130]}'

def storefront_asins(html):
    s=BeautifulSoup(html,'lxml');result=[]
    for el in s.select('[data-asin]'):
        asin=(el.get('data-asin') or '').upper()
        if re.fullmatch(r'[A-Z0-9]{10}',asin) and asin not in result:result.append(asin)
    return result

def collect_storefront(page,seller_id,market,max_pages=3):
    collected=[]; seen=set();log=[]
    for index in range(1,max_pages+1):
        url=url_store(seller_id,market)+f'&page={index}'
        html,state=navigate(page,url)
        log.append({'url':url,'status':state,'observed_at':datetime.now(timezone.utc).isoformat(timespec='seconds')})
        if not html:break
        page_asins=storefront_asins(html)
        novel=[a for a in page_asins if a not in seen]
        if not novel:break
        seen.update(novel);collected.extend(novel)
    return collected,log

PRODUCT_COLS=[('marketplace','Marketplace'),('asin','ASIN'),('product_url','Product URL'),('image_url','Image URL'),
('title','Product Title'),('highlights','Item Highlights'),('device_model','Device Model'),('price_text','Price Raw'),
('price_numeric','Price Number'),('currency','Currency'),('buybox','Buy Box'),('in_stock','In Stock'),
('ships_from','Ships From Raw'),('ships_from_entity','Ships From Entity'),('sold_by','Sold By'),('seller_id','Seller ID'),('variations','Variations'),
('offer_count','Offer Count'),('rating','Product Rating'),('rating_count','Review Count'),('bsr','Best Sellers Rank'),
('date_first_available','Date First Available'),('observed_at','Observed UTC'),('status','Data Status'),
('delivery_zip','Delivery ZIP / City')]
SELLER_COLS=['seller_id','marketplace','name','seller_url','storefront_url','legal_name','business_address',
'country_code','positive_12m_percent','positive_12m_text','visible_rating','visible_rating_count',
'contact_email','contact_phone','company_registration_number','vat_number','eori_number',
'registration_documents','contact_status','registration_status','return_policy','page_status','observed_at','status','feedback_period','profile_source']

def safe_filename(s):return re.sub(r'[^A-Za-z0-9_.-]','_',s)[:85]

def style_sheet(sheet, ncols, default_width=20):
    for c in range(1,ncols+1):
        sheet.column_dimensions[get_column_letter(c)].width=default_width
        cell=sheet.cell(row=1,column=c)
        cell.fill=PatternFill('solid',fgColor='183153');cell.font=Font(color='FFFFFF',bold=True)
        cell.alignment=Alignment(wrap_text=True,vertical='center')
    sheet.row_dimensions[1].height=35
    sheet.freeze_panes='A2'

ILLEGAL_RE=re.compile(r'[\x00-\x08\x0b\x0c\x0e-\x1f]')
def cell_value(v):
    if isinstance(v,(dict,list)):v=json.dumps(v,ensure_ascii=False)
    if isinstance(v,str):
        v=ILLEGAL_RE.sub('',v)
        # Values from exports/pages are data, never executable Excel formulas.
        return "'"+v if v.startswith('=') else v
    return v

def export_report(out_path,seller,products,events):
    wb=Workbook()
    dashboard=wb.active;dashboard.title='Dashboard'
    for row in [['SELLER INTELLIGENCE','VALUE'],['Seller ID',seller['seller_id']],
       ['Marketplace',seller['marketplace']],['Seller',seller.get('name','')],['Legal Country',seller.get('country_code','')],
       ['Observed products',len(products)],['Product rating < 4.1',sum(isinstance(p.get('rating'),(float,int)) and p['rating']<4.1 for p in products)],
       ['Stock YES',sum(p.get('in_stock')=='YES' for p in products)],['Catalog completeness','OBSERVED ONLY, NOT VERIFIED FULL'],
       ['Collected UTC',datetime.now(timezone.utc).isoformat(timespec='seconds')]]:
        dashboard.append([cell_value(v) for v in row])
    style_sheet(dashboard,2);dashboard.column_dimensions['A'].width=35;dashboard.column_dimensions['B'].width=48
    sht=wb.create_sheet('Products');heads=[c[1] for c in PRODUCT_COLS]
    sht.append(heads)
    for p in products:sht.append([cell_value(p.get(k)) for k,_ in PRODUCT_COLS])
    style_sheet(sht,len(heads))
    for col,width in {'C':44,'D':44,'E':54,'F':65,'G':27,'M':35,'T':38}.items():sht.column_dimensions[col].width=width
    if products:
        # Highlight whole product row whenever a numeric rating is known to be below 4.1.
        rating_col=get_column_letter([k for k,_ in PRODUCT_COLS].index('rating')+1)
        sht.conditional_formatting.add(f'A2:{get_column_letter(len(heads))}{len(products)+1}',
            FormulaRule(formula=[f'AND(ISNUMBER(${rating_col}2),${rating_col}2<4.1)'],
                        fill=PatternFill('solid',bgColor='FEE2E2',fgColor='FEE2E2'),font=Font(color='991B1B')))
    profile=wb.create_sheet('Seller Profile')
    profile.append(['Seller property','Value'])
    for k in SELLER_COLS:profile.append([k,cell_value(seller.get(k))])
    for period in ('1 month','3 months','12 months','lifetime'):
        d=seller.get('feedback',{}).get(period,{})
        for fieldn in ('rating','ratings','positive_percent','status'):
            profile.append([f'Feedback {period} / {fieldn}',cell_value(d.get(fieldn))])
    style_sheet(profile,2);profile.column_dimensions['A'].width=38;profile.column_dimensions['B'].width=98
    details=wb.create_sheet('Item Details');details.append(['ASIN','Field','Value'])
    for p in products:
        for k,v in p.get('item_details',{}).items():details.append([p['asin'],cell_value(k),cell_value(v)])
    style_sheet(details,3);details.column_dimensions['B'].width=40;details.column_dimensions['C'].width=90
    log=wb.create_sheet('Run Log');log.append(['UTC','URL','Status'])
    for e in events:log.append([e.get('observed_at',''),e.get('url',''),cell_value(e.get('status',''))])
    style_sheet(log,3);log.column_dimensions['B'].width=80;log.column_dimensions['C'].width=52
    wb.save(str(out_path))

SEED_COLS=('asin','url','product url','product_url','link','amazon url')

def load_seeds(file):
    path=Path(file)
    if path.suffix.lower() in ('.csv','.tsv'):
        with path.open(encoding='utf-8-sig',newline='') as f:
            reader=csv.DictReader(f,delimiter='\t' if path.suffix.lower()=='.tsv' else ',')
            return list(reader)
    if path.suffix.lower()=='.json':
        val=json.loads(path.read_text(encoding='utf-8'))
        return val if isinstance(val,list) else val.get('items',[])
    if path.suffix.lower()=='.xlsx':
        wb=load_workbook(str(path),data_only=True)
        result=[]; found_header=False
        try:
            for ws in wb.worksheets:
                rows=ws.iter_rows(); headers=None
                for n,r in enumerate(rows):
                    if n >= 10:break
                    candidate=[clean(str(c.value or '')).lower() for c in r]
                    if any(h in SEED_COLS for h in candidate):
                        headers=candidate;found_header=True;break
                if headers is None:continue
                for r in rows:
                    row={}
                    for h,c in zip(headers,r):
                        if not h:continue
                        v=str(c.value).strip() if c.value is not None else ''
                        target=c.hyperlink.target if c.hyperlink else None
                        if h in SEED_COLS and target and market_from_url(target,None):v=target
                        row[h]=v
                    if any(row.values()):result.append(row)
        finally:
            wb.close()
        if found_header:return result
        raise ValueError('No ASIN or URL column found in XLSX.')
    raise ValueError('Supported input: CSV, TSV, JSON, XLSX')

def market_from_name(name):
    low=name.lower()
    for market,(domain,_) in sorted(MARKETS.items(),key=lambda item:len(item[1][0]),reverse=True):
        if domain.removeprefix('www.') in low:return market
    for tok in re.split(r'[^a-z0-9]+',low):
        t={'gb':'uk'}.get(tok,tok).upper()
        if len(tok)==2 and t in MARKETS:return t
    return None

def load_folder(folder,default_market=None):
    """Read Xray files deterministically, without interactive prompts (also safe in CI)."""
    folder=Path(folder)
    files=sorted(f for f in folder.iterdir() if f.is_file() and f.suffix.lower()=='.xlsx' and not f.name.startswith('~$'))
    if not files:raise ValueError(f'No .xlsx files in {folder}. Add a Helium 10 Xray export first.')
    seeds=[]
    for f in files:
        market=market_from_name(f.stem) or default_market
        try:
            found=resolve_seeds(load_seeds(f),market)
            if not found:raise ValueError('No valid ASIN seeds found')
        except Exception as exc:raise ValueError(f'{f.name}: {exc}') from exc
        LOG.info('%s: %d ASIN(s), default marketplace %s',f.name,len(found),market)
        seeds.extend(found)
    return list(dict.fromkeys(seeds))

def resolve_seeds(rows,default_market):
    result=[]
    for row in rows:
        if not isinstance(row,dict):raise ValueError('Every input row must be an object with ASIN or URL')
        row={clean(str(k)).lower():v for k,v in row.items() if k}
        values=[str(row[k]).strip() for k in SEED_COLS if row.get(k)]
        asin=next((asin_from(v) for v in values if asin_from(v)),None)
        if not asin:
            LOG.warning('Could not resolve ASIN from %r',values);continue
        url_market=next((market_from_url(v,None) for v in values if asin_from(v)==asin and market_from_url(v,None)),None)
        market=clean(str(row.get('marketplace') or row.get('market') or url_market or default_market or '')).upper()
        market={'GB':'UK'}.get(market,market)
        if market not in MARKETS:
            raise ValueError(f'Marketplace missing or invalid for {asin}: {market!r}. Use --marketplace US (or another code), a Marketplace column, an Amazon URL, or a file name such as Xray_US.xlsx.')
        result.append((asin,market))
    return list(dict.fromkeys(result))

def _set_delivery_once(page,market,zipcode):
    """Set the ordinary 'Deliver to' postal code via the site's own UI. Returns (status, shown_location)."""
    if not zipcode:return 'NO_ZIP_FOR_MARKET',''
    html,state=navigate(page,f'https://{MARKETS[market][0]}/')
    if not html:return state,''
    try:
        try:page.wait_for_load_state('load',timeout=15000)
        except Exception:pass
        probe={'SA':'#GLUXCityList','AE':'#GLUXCityWithDistrictCityInput'}.get(market,'#GLUXZipUpdateInput')
        for _ in range(3):  # the header link is sometimes clicked before its handler is attached
            page.locator('#nav-global-location-popover-link, #glow-ingress-block').first.click(timeout=8000)
            try:
                page.locator(probe).first.wait_for(state='attached' if market=='SA' else 'visible',timeout=5000);break
            except Exception:page.wait_for_timeout(1000)
        if market=='SA':
            page.locator('#GLUXCityList').first.wait_for(state='attached',timeout=8000)
            page.locator('#GLUXCityList').first.select_option(label=zipcode)
        elif market=='AE':
            city=page.locator('#GLUXCityWithDistrictCityInput').first
            city.wait_for(timeout=8000);city.click();city.type(zipcode,delay=80);page.wait_for_timeout(1200)
            page.locator('#GLUXCityWithDistrictCityList li',has_text=zipcode).first.click(timeout=5000)
            page.locator('#GLUXCityWithDistrictApplyButton').first.click(timeout=5000)
        else:
            inp=page.locator('#GLUXZipUpdateInput').first
            inp.wait_for(timeout=8000)
            parts=zipcode.split(' ')
            if page.locator('#GLUXZipUpdateInput_1').count() and len(parts)==2:
                inp.fill(parts[0]);page.locator('#GLUXZipUpdateInput_1').first.fill(parts[1]);inp.press('Enter')
            else:
                inp.click();inp.type(zipcode,delay=60);inp.press('Enter')
            page.wait_for_timeout(2000)
            if inp.is_visible():  # Enter did not submit: use the Apply button
                page.locator('#GLUXZipUpdate input, #GLUXZipUpdate').first.click(timeout=5000)
        page.wait_for_timeout(1500)
        for sel in ('#GLUXConfirmClose','button[name="glowDoneButton"]','.a-popover-footer .a-button-primary input'):
            try:
                b=page.locator(sel).first
                if b.count() and b.is_visible():b.click(timeout=3000);break
            except Exception:continue
        page.wait_for_timeout(1200)
        page.reload(wait_until='domcontentloaded');page.wait_for_timeout(1000)
        shown=clean(page.locator('#glow-ingress-line2').first.inner_text(timeout=3000))
        ok=zipcode.replace(' ','').lower() in shown.replace(' ','').lower()
        return ('OK' if ok else 'UNVERIFIED'),shown
    except Exception as exc:
        if browser_closed(exc) or isinstance(exc,BrowserClosedError):raise BrowserClosedError(str(exc)) from exc
        return f'ERROR:{type(exc).__name__}:{str(exc)[:100]}',''

def set_delivery(page,market,zipcode,attempts=3):
    status,shown='ERROR',''
    for _ in range(attempts):
        status,shown=_set_delivery_once(page,market,zipcode)
        if status in ('OK','NO_ZIP_FOR_MARKET','BLOCKED_CHALLENGE'):break
        page.wait_for_timeout(1500)
    return status,shown

def export_run_report(out_path,seeds,products,groups,events,run_status='UNKNOWN'):
    """Export an overall result, including unavailable seed products."""
    wb=Workbook(); summary=wb.active;summary.title='Summary'
    for row in [
        ['Run summary','Value'],['Seed ASINs',len(seeds)],
        ['Seller profiles collected',sum((g.get('seller') or {}).get('status')=='OK' for g in groups.values())],
        ['Observed / unavailable products',len(products)],
        ['Products with status OK',sum(p.get('status')=='OK' for p in products)],
        ['Products requiring attention',sum(p.get('status')!='OK' for p in products)],
        ['Collected UTC',datetime.now(timezone.utc).isoformat(timespec='seconds')],
        ['Coverage','Observed products only; unavailable data remains blank'],
        ['Run status',run_status],
        ['Sellers identified',len(groups)],
        ['Seed products not processed',sum(p.get('status')=='NOT_PROCESSED' for p in products)],
    ]:summary.append(row)
    style_sheet(summary,2);summary.column_dimensions['A'].width=40;summary.column_dimensions['B'].width=70
    sheet=wb.create_sheet('Products');sheet.append([label for _,label in PRODUCT_COLS])
    for p in products:sheet.append([cell_value(p.get(k)) for k,_ in PRODUCT_COLS])
    style_sheet(sheet,len(PRODUCT_COLS));sheet.column_dimensions['E'].width=54
    sellers=wb.create_sheet('Sellers');sellers.append(SELLER_COLS)
    feedback=wb.create_sheet('Seller Feedback')
    feedback.append(['Seller ID','Marketplace','Period','Rating','Rating Count','Positive %','Status'])
    for group in groups.values():
        seller=group.get('seller')
        if not seller:
            seller=parse_seller('',group['seller_id'],group['marketplace'])
            seller['status']=seller['page_status']='NOT_PROCESSED'
        sellers.append([cell_value(seller.get(k)) for k in SELLER_COLS])
        for period in ('1 month','3 months','12 months','lifetime'):
            f=seller.get('feedback',{}).get(period,{})
            feedback.append([seller['seller_id'],seller['marketplace'],period,f.get('rating'),f.get('ratings'),f.get('positive_percent'),f.get('status','NOT_COLLECTED')])
    style_sheet(sellers,len(SELLER_COLS))
    style_sheet(feedback,7);feedback.auto_filter.ref=feedback.dimensions
    inputs=wb.create_sheet('Input Seeds');inputs.append(['ASIN','Marketplace','Product URL'])
    for asin,market in seeds:inputs.append([asin,market,url_asin(asin,market)])
    style_sheet(inputs,3);inputs.column_dimensions['C'].width=65
    log=wb.create_sheet('Run Log');log.append(['UTC','URL','Status'])
    for event in events:log.append([cell_value(event.get(k,'')) for k in ('observed_at','url','status')])
    style_sheet(log,3);log.column_dimensions['B'].width=80;log.column_dimensions['C'].width=70
    for ws in (sheet,sellers,inputs,log):ws.auto_filter.ref=ws.dimensions
    out_path=Path(out_path)
    with tempfile.NamedTemporaryFile(dir=out_path.parent,suffix='.xlsx',delete=False) as tmp:
        tmp_path=Path(tmp.name)
    try:
        wb.save(tmp_path);tmp_path.replace(out_path)
    finally:
        wb.close();tmp_path.unlink(missing_ok=True)

def missing_product(asin,market,status):
    return {'asin':asin,'marketplace':market,'product_url':url_asin(asin,market),
            'status':status,'observed_at':datetime.now(timezone.utc).isoformat(timespec='seconds')}

def write_json_atomic(path, payload):
    path=Path(path);tmp=path.with_suffix(path.suffix+'.tmp')
    tmp.write_text(json.dumps(payload,ensure_ascii=False,indent=2),encoding='utf-8')
    tmp.replace(path)

def load_resume(folder):
    folder=Path(folder)
    if folder.is_file():folder=folder.parent
    checkpoint=folder/'checkpoint.json'
    if checkpoint.exists():
        state=json.loads(checkpoint.read_text(encoding='utf-8'))
        if state.get('version')!=1:raise ValueError('Unsupported checkpoint version')
        return state
    # Import older reports: keep collected product data but retry failed profiles.
    result=folder/'result.xlsx'
    if not result.exists():raise ValueError(f'No result.xlsx or checkpoint.json in {folder}')
    wb=load_workbook(result,read_only=True,data_only=True)
    try:
        seeds=[(str(r[0]),str(r[1])) for r in list(wb['Input Seeds'].values)[1:] if r[0] and r[1]]
        rows=iter(wb['Products'].values);headers=next(rows)
        keys={label:key for key,label in PRODUCT_COLS}
        products=[{keys[h]:v for h,v in zip(headers,row) if h in keys} for row in rows]
    finally:wb.close()
    groups=[]
    for path in sorted(folder.glob('*.json')):
        data=json.loads(path.read_text(encoding='utf-8'))
        seller=data.get('seller') or {}
        if not seller.get('seller_id'):continue
        good=seller.get('name') and seller.get('page_status','HTTP_200')=='HTTP_200'
        groups.append({'marketplace':seller['marketplace'],'seller_id':seller['seller_id'],
                       'seller':seller if good else None,'products':{r['asin']:r for r in data.get('products',[])},
                       'events':[],'catalog_done':False})
        products.extend(data.get('products',[]))
    return {'version':1,'seeds':seeds,'products':products,'groups':groups,'events':[]}

def run(seeds,out_dir,headless=True,max_pages=3,max_products=100,delay=2.0,zips=None,resume_data=None,skip_storefront=False):
    out_dir=Path(out_dir);out_dir.mkdir(parents=True,exist_ok=True)
    old=resume_data or {};selected={(m,a) for a,m in seeds}
    products={(p['marketplace'],p['asin']):p for p in old.get('products',[]) if (p['marketplace'],p['asin']) in selected}
    seller_keys={(p['marketplace'],p['seller_id']) for p in products.values() if p.get('seller_id')}
    groups={(g['marketplace'],g['seller_id']):g for g in old.get('groups',[]) if (g['marketplace'],g['seller_id']) in seller_keys}
    events=[];run_status='RUNNING';located={};current_market=None
    zips={**ZIPS,**CITY_MARKETS,**(zips or {})}
    def add_event(url,status,sink=events):
        sink.append({'url':url,'status':status,'observed_at':datetime.now(timezone.utc).isoformat(timespec='seconds')})
    def checkpoint(export=False):
        all_products=dict(products)
        for group in groups.values():
            for row in group['products'].values():all_products.setdefault((row['marketplace'],row['asin']),row)
        for asin,market in seeds:all_products.setdefault((market,asin),missing_product(asin,market,'NOT_PROCESSED'))
        logs=events+[e for g in groups.values() for e in g['events']]
        write_json_atomic(out_dir/'checkpoint.json',{'version':1,'run_status':run_status,'seeds':seeds,
            'products':list(all_products.values()),'groups':list(groups.values()),'events':logs})
        if export:export_run_report(out_dir/'result.xlsx',seeds,list(all_products.values()),groups,logs,run_status)
    def save_group(market,seller_id,g):
        if not g.get('seller'):return
        filename=out_dir/f'{market}_{safe_filename(seller_id)}'
        export_report(filename.with_suffix('.xlsx'),g['seller'],list(g['products'].values()),g['events'])
        write_json_atomic(filename.with_suffix('.json'),{'seller':g['seller'],'products':list(g['products'].values()),'events':g['events']})
    def ensure_location(page,market):
        nonlocal current_market
        if current_market==market:return
        status,shown=set_delivery(page,market,zips.get(market))
        located[market]=zips.get(market,'') if status=='OK' else ''
        current_market=market
        add_event(f'delivery:{market}:{zips.get(market,"")}',f'{status} {shown}'.strip())
        if status!='OK':LOG.warning('Delivery location for %s not confirmed: %s',market,status)
    def collect_profile(page,market,seller_id,g):
        previous=g.get('seller') or {}
        if previous.get('name') and previous.get('status') in ('OK','RETAIL_PROFILE_NOT_APPLICABLE'):return
        product=next(iter(g['products'].values()))
        name=product.get('sold_by','')
        if name.lower() in ('amazon',MARKETS[market][0].removeprefix('www.')):
            seller=parse_seller('',seller_id,market)
            seller.update(name=name,status='RETAIL_PROFILE_NOT_APPLICABLE',page_status='NOT_APPLICABLE',profile_source=product['product_url'])
            g['seller']=seller
            add_event(product['product_url'],'RETAIL_PROFILE_NOT_APPLICABLE: Sold by '+name,g['events'])
            save_group(market,seller_id,g);checkpoint(export=True)
            return
        url=url_seller(seller_id,market);html,state=navigate(page,url)
        add_event(url,state,g['events'])
        seller=parse_seller(html,seller_id,market);seller['page_status']=state;seller['profile_source']=url
        if not html:
            seller['status']=state
            seller['contact_status']=seller['registration_status']='NOT_COLLECTED'
        g['seller']=seller
        save_group(market,seller_id,g);checkpoint(export=True)
        LOG.info('Seller %s: %s / %s',seller_id,seller.get('name') or '(not retrieved)',seller['status'])
        time.sleep(delay)
    try:
        checkpoint(export=True)
        with sync_playwright() as pw:
            browser=pw.chromium.launch(headless=headless)
            context=browser.new_context(locale='en-US',viewport={'width':1450,'height':1050})
            page=context.new_page()
            try:
                for index,(asin,market) in enumerate(seeds,1):
                    LOG.info('Product %d/%d: %s %s',index,len(seeds),market,asin)
                    ensure_location(page,market)
                    product=products.get((market,asin))
                    if not (product and product.get('title') and product.get('seller_id') and product.get('status')=='OK'):
                        url=url_asin(asin,market);html,state=navigate(page,url);add_event(url,state)
                        product=parse_product(html,asin,market) if html else missing_product(asin,market,state)
                        product['delivery_zip']=located.get(market,'')
                        if html and not product.get('seller_id'):product['status']='NO_SELLER_ID' if product.get('title') else 'UNVERIFIED'
                        products[(market,asin)]=product
                        checkpoint()
                        time.sleep(delay)
                    if product.get('seller_id'):
                        key=(market,product['seller_id'])
                        g=groups.setdefault(key,{'marketplace':market,'seller_id':key[1],'products':{},'seller':None,'events':[],'catalog_done':False})
                        g['products'][asin]=product
                        # Fetch the profile now, not after hundreds of other listings.
                        collect_profile(page,market,key[1],g)
                    checkpoint(export=index%10==0)
                if not skip_storefront:
                    for (market,seller_id),g in groups.items():
                        if g.get('catalog_done'):continue
                        ensure_location(page,market)
                        catalog,logs=collect_storefront(page,seller_id,market,max_pages);g['events'].extend(logs)
                        for asin in catalog[:max_products]:
                            if asin in g['products']:continue
                            url=url_asin(asin,market);html,state=navigate(page,url);add_event(url,state,g['events'])
                            p=parse_product(html,asin,market) if html else missing_product(asin,market,state)
                            p['delivery_zip']=located.get(market,'')
                            if p.get('seller_id') and p['seller_id']!=seller_id:p['status']='DIFFERENT_BUYBOX_SELLER'
                            g['products'][asin]=p
                            checkpoint();time.sleep(delay)
                        g['catalog_done']=bool(logs) and all(e['status']=='HTTP_200' for e in logs)
                        save_group(market,seller_id,g);checkpoint(export=True)
                good_products=all(products.get((m,a),{}).get('status')=='OK' for a,m in seeds)
                good_sellers=bool(groups) and all((g.get('seller') or {}).get('status') in ('OK','RETAIL_PROFILE_NOT_APPLICABLE') for g in groups.values())
                good_catalog=skip_storefront or all(g.get('catalog_done') for g in groups.values())
                run_status='COMPLETED' if good_products and good_sellers and good_catalog else 'PARTIAL'
            finally:
                if browser.is_connected():browser.close()
    except (BrowserClosedError,KeyboardInterrupt) as exc:
        run_status='INTERRUPTED';add_event('run',f'{type(exc).__name__}: {exc}')
        raise
    except Exception as exc:
        run_status='FAILED';add_event('run',f'ERROR:{type(exc).__name__}:{exc}')
        raise
    finally:
        # Persist recoverable data before attempting any formatted exports.
        checkpoint()
        for (market,seller_id),g in groups.items():save_group(market,seller_id,g)
        checkpoint(export=True)
        LOG.info('Run %s. Result: %s. Resume with --resume "%s"',run_status,out_dir/'result.xlsx',out_dir)
    return sum((g.get('seller') or {}).get('status')=='OK' for g in groups.values())

def load_config(path):
    defaults={'input_dir':'xray helium10','output_dir':'reports','marketplace':None,
              'headless':True,'max_pages':3,'max_products':100,'delay':2.0}
    raw=json.loads(Path(path).read_text(encoding='utf-8-sig'))
    if not isinstance(raw,dict):raise ValueError('Config must contain a JSON object')
    unknown=set(raw)-set(defaults)
    if unknown:raise ValueError(f'Unknown config keys: {sorted(unknown)}')
    defaults.update(raw)
    for key in ('input_dir','output_dir'):
        if not isinstance(defaults[key],str) or not defaults[key].strip():raise ValueError(f'{key} must be a nonempty path')
    if defaults['marketplace'] is not None:
        defaults['marketplace']=str(defaults['marketplace']).strip().upper()
        if defaults['marketplace'] not in MARKETS:raise ValueError('Invalid config marketplace')
    if not isinstance(defaults['headless'],bool):raise ValueError('headless must be true or false')
    for key in ('max_pages','max_products'):
        if type(defaults[key]) is not int or defaults[key]<1:raise ValueError(f'{key} must be a positive integer')
    if type(defaults['delay']) not in (float,int) or not math.isfinite(defaults['delay']) or defaults['delay']<0:
        raise ValueError('delay must be a finite nonnegative number')
    return defaults

def project_path(value):
    path=Path(value).expanduser()
    return path if path.is_absolute() else BASE_DIR/path

def main(argv=None):
    ap=argparse.ArgumentParser(description=__doc__)
    source=ap.add_mutually_exclusive_group()
    source.add_argument('--input',help='Xray folder or one CSV/TSV/JSON/XLSX file')
    source.add_argument('--asin',help='Seed ASIN; requires marketplace')
    source.add_argument('--url',help='Amazon product URL')
    source.add_argument('--resume',help='Resume a run directory or its result.xlsx, including older reports')
    ap.add_argument('--config',default=str(BASE_DIR/'config.json'),help='JSON config file')
    ap.add_argument('--marketplace',type=str.upper,choices=sorted(MARKETS),help='Fallback market when row, URL and file name do not identify one')
    ap.add_argument('--out',help='Output root; each run creates its own timestamped folder')
    ap.add_argument('--headless',action=argparse.BooleanOptionalAction,default=None,help='No browser window; automatically enabled in CI')
    ap.add_argument('--max-pages',type=int)
    ap.add_argument('--max-products',type=int)
    ap.add_argument('--delay',type=float)
    ap.add_argument('--validate-only',action='store_true',help='Validate input and list ASINs without visiting Amazon')
    ap.add_argument('--limit',type=int,help='Process only the first N input seeds (for a test run)')
    ap.add_argument('--skip-storefront',action='store_true',help='Collect input listings and seller profiles without expanding storefront catalogs')
    ap.add_argument('--zip',action='append',default=[],metavar='MARKET=CODE',help='Override delivery postal code, e.g. --zip US=90210 (repeatable)')
    args=ap.parse_args(argv);logging.basicConfig(level=logging.INFO,format='%(levelname)s %(message)s')
    try:
        config=load_config(project_path(args.config))
        for key in ('marketplace','headless','max_pages','max_products','delay'):
            if getattr(args,key) is None:setattr(args,key,config[key])
        if os.environ.get('CI','').lower() in ('1','true','yes'):args.headless=True
        if args.max_pages<1 or args.max_products<1 or not math.isfinite(args.delay) or args.delay<0:
            raise ValueError('max-pages and max-products must be positive; delay must be finite and nonnegative')
        resume_data=None
        if args.limit is not None and args.limit<1:raise ValueError('--limit must be positive')
        if args.resume:
            resume_data=load_resume(project_path(args.resume));seeds=[tuple(s) for s in resume_data['seeds']]
        elif args.asin or args.url:
            seeds=resolve_seeds([{'asin':args.asin} if args.asin else {'url':args.url}],args.marketplace)
        else:
            src=project_path(args.input or config['input_dir'])
            if not src.exists():raise ValueError(f'Input not found: {src}')
            seeds=load_folder(src,args.marketplace) if src.is_dir() else resolve_seeds(load_seeds(src),market_from_name(src.stem) or args.marketplace)
        if not seeds:raise ValueError('No valid ASIN seeds found')
        if args.limit is not None:seeds=seeds[:args.limit]
        zips={}
        for override in args.zip:
            key,sep,value=override.partition('=');key=key.strip().upper();value=value.strip()
            if not sep or key not in MARKETS or not value:raise ValueError('--zip must be MARKET=CODE, e.g. US=90210')
            zips[key]=value
    except (ValueError,OSError) as exc:ap.error(str(exc))
    if args.validate_only:
        print(json.dumps({'seed_count':len(seeds),'seeds':[{'asin':a,'marketplace':m} for a,m in seeds]},ensure_ascii=False,indent=2))
        return 0
    out=project_path(args.out or config['output_dir'])/datetime.now(timezone.utc).strftime('%Y%m%d_%H%M%S_%f_UTC')
    out.mkdir(parents=True,exist_ok=True)
    file_log=logging.FileHandler(out/'run.log',encoding='utf-8')
    file_log.setFormatter(logging.Formatter('%(asctime)s %(levelname)s %(message)s'))
    logging.getLogger().addHandler(file_log)
    try:
        n=run(seeds,out,args.headless,args.max_pages,args.max_products,args.delay,zips,resume_data,args.skip_storefront)
        state=json.loads((out/'checkpoint.json').read_text(encoding='utf-8'))
        print(f'{state["run_status"]}: {n} seller profile(s). Result: {out / "result.xlsx"}')
        if state['run_status']!='COMPLETED':LOG.warning('Run is incomplete. See result.xlsx / Run Log; resume with --resume "%s".',out)
        return 0 if state['run_status']=='COMPLETED' else 2
    except KeyboardInterrupt:
        LOG.warning('Interrupted. Resume with --resume "%s"',out)
        return 130
    except Exception:
        LOG.exception('Run failed. Diagnostic result, if saved: %s',out/'result.xlsx')
        return 1
    finally:
        logging.getLogger().removeHandler(file_log);file_log.close()
if __name__=='__main__':sys.exit(main())
