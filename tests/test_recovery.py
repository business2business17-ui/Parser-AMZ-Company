import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import MagicMock, patch

from openpyxl import Workbook, load_workbook

import agent


SELLER_HTML = '''<h1 id="seller-name">Example Seller</h1>
<div id="seller-info-feedback-summary">5 out of 5 stars 98% positive in the last 12 months (600 ratings)</div>
<div id="page-section-detail-seller-info"><div>Business Name: Example GmbH</div>
<div>Business Address: Teststrasse 1 Berlin DE</div></div>
<div id="rating-thirty"><span class="ratings-reviews">4.7</span><span class="ratings-reviews-count">34</span></div>
<div id="rating-ninety"><span class="ratings-reviews">4.8</span><span class="ratings-reviews-count">111</span></div>
<div id="rating-year"><span class="ratings-reviews">4.9</span><span class="ratings-reviews-count">600</span></div>
<div id="rating-lifetime"><span class="ratings-reviews">4.6</span><span class="ratings-reviews-count">2,563</span></div>
<footer>Copyright 1996-2026 <a href="/ap/register">Register</a><a href="/business">Amazon Business</a>
<a href="mailto:amazon@example.com">Support</a> Phone: +1 999 999 9999</footer>'''
PRODUCT_HTML = '''<span id="productTitle">Example product</span>
<div id="merchantInfoFeature_feature_div"><a id="sellerProfileTriggerId" href="/sp?seller=A123456789">Example Seller</a></div>
<div id="fulfillerInfoFeature_feature_div"><span class="offer-display-feature-text-message">Amazon</span></div>
<div id="acrPopover"><span class="a-icon-alt">4.6 out of 5 stars</span></div><span id="acrCustomerReviewText">321 ratings</span>'''


class SellerParsingTests(unittest.TestCase):
    def test_exact_rating_and_separate_periods(self):
        s=agent.parse_seller(SELLER_HTML,'A123456789','DE')
        self.assertEqual((s['visible_rating'],s['visible_rating_count']),(4.9,600))
        self.assertEqual(s['feedback']['1 month']['rating'],4.7)
        self.assertEqual(s['feedback']['3 months']['ratings'],111)
        self.assertEqual(s['feedback']['lifetime']['ratings'],2563)
        self.assertEqual(s['positive_12m_percent'],98)

    def test_footer_is_not_seller_contact_or_registration(self):
        s=agent.parse_seller(SELLER_HTML,'A123456789','DE')
        for key in ('contact_phone','contact_email','registration_documents','company_registration_number'):
            self.assertEqual(s[key],'')
        self.assertEqual(s['contact_status'],'NOT_SHOWN')
        self.assertEqual(s['registration_status'],'NOT_SHOWN')
        self.assertEqual(s['legal_name'],'Example GmbH')
        self.assertEqual(s['country_code'],'DE')

    def test_registration_fields_do_not_bleed_into_each_other(self):
        html='''<h1 id="seller-name">Test</h1><div id="page-section-detail-seller-info">
        Business Name: Test Ltd Business Address: Road 1 GB Phone: +44 20 1234 5678
        Email: contact@example.com Company registration number: 01234567 VAT ID: GB123456789 EORI: GB987654321
        </div>'''
        s=agent.parse_seller(html,'A123456789','UK')
        self.assertEqual(s['contact_phone'],'+44 20 1234 5678')
        self.assertEqual(s['contact_email'],'contact@example.com')
        self.assertEqual(s['company_registration_number'],'01234567')
        self.assertEqual(s['vat_number'],'GB123456789')
        self.assertEqual(s['eori_number'],'GB987654321')

    def test_unvisited_seller_is_not_marked_not_shown(self):
        s=agent.parse_seller('','A123456789','US')
        self.assertEqual(s['contact_status'],'NOT_COLLECTED')
        self.assertEqual(s['registration_status'],'NOT_COLLECTED')

    def test_product_uses_buybox_seller_not_recommendations(self):
        p=agent.parse_product('<a href="/sp?seller=WRONGSELLER">Sponsored</a>'+PRODUCT_HTML,'B000000001','US')
        self.assertEqual(p['seller_id'],'A123456789')
        self.assertEqual(p['ships_from'],'Amazon')
        self.assertEqual(p['ships_from_entity'],'AMAZON')
        self.assertEqual(p['rating_count'],321)


class RecoveryTests(unittest.TestCase):
    def test_closed_page_is_fatal_not_one_error_per_remaining_url(self):
        page=MagicMock()
        page.goto.side_effect=RuntimeError('Target page, context or browser has been closed')
        with self.assertRaises(agent.BrowserClosedError):agent.navigate(page,'https://www.amazon.com/')
        self.assertEqual(page.goto.call_count,1)

    def test_interrupt_writes_checkpoint_and_remaining_rows(self):
        with tempfile.TemporaryDirectory() as tmp:
            with patch.object(agent,'sync_playwright'), patch.object(agent,'set_delivery',return_value=('OK','10001')), patch.object(agent,'navigate',side_effect=agent.BrowserClosedError('closed')) as nav:
                with self.assertRaises(agent.BrowserClosedError):
                    agent.run([('B000000001','US'),('B000000002','US')],tmp,delay=0)
            self.assertEqual(nav.call_count,1)
            state=json.loads((Path(tmp)/'checkpoint.json').read_text(encoding='utf-8'))
            self.assertEqual(state['run_status'],'INTERRUPTED')
            self.assertEqual([p['status'] for p in state['products']],['NOT_PROCESSED','NOT_PROCESSED'])
            self.assertTrue((Path(tmp)/'result.xlsx').exists())

    def test_profile_is_collected_before_next_product_and_survives_interrupt(self):
        with tempfile.TemporaryDirectory() as tmp:
            with patch.object(agent,'sync_playwright'), patch.object(agent,'set_delivery',return_value=('OK','10001')), patch.object(agent,'navigate',side_effect=[(PRODUCT_HTML,'HTTP_200'),(SELLER_HTML,'HTTP_200'),agent.BrowserClosedError('closed')]) as nav:
                with self.assertRaises(agent.BrowserClosedError):
                    agent.run([('B000000001','US'),('B000000002','US')],tmp,delay=0,skip_storefront=True)
            self.assertIn('/sp?',nav.call_args_list[1].args[1])
            state=agent.load_resume(tmp)
            self.assertEqual(state['groups'][0]['seller']['name'],'Example Seller')
            self.assertEqual(state['run_status'],'INTERRUPTED')
            with patch.object(agent,'sync_playwright'), patch.object(agent,'set_delivery',return_value=('OK','10001')), patch.object(agent,'navigate',return_value=(PRODUCT_HTML,'HTTP_200')) as nav:
                n=agent.run([('B000000001','US'),('B000000002','US')],tmp,delay=0,skip_storefront=True,resume_data=state)
            self.assertEqual(n,1)
            self.assertEqual(nav.call_count,1)
            self.assertIn('B000000002',nav.call_args.args[1])
            self.assertEqual(agent.load_resume(tmp)['run_status'],'COMPLETED')
            wb=load_workbook(Path(tmp)/'result.xlsx',data_only=True)
            self.assertEqual(wb['Products'].max_row,3)
            feedback=list(wb['Seller Feedback'].values)
            self.assertEqual(feedback[3][3:5],(4.9,600))
            wb.close()

    def test_failed_profile_does_not_count_as_completed(self):
        with tempfile.TemporaryDirectory() as tmp:
            with patch.object(agent,'sync_playwright'), patch.object(agent,'set_delivery',return_value=('OK','10001')), patch.object(agent,'navigate',side_effect=[(PRODUCT_HTML,'HTTP_200'),('','BLOCKED_CHALLENGE')]):
                n=agent.run([('B000000001','US')],tmp,delay=0,skip_storefront=True)
            state=agent.load_resume(tmp)
            self.assertEqual(n,0)
            self.assertEqual(state['run_status'],'PARTIAL')
            self.assertEqual(state['groups'][0]['seller']['contact_status'],'NOT_COLLECTED')

    def test_legacy_report_resume_retries_missing_profiles(self):
        with tempfile.TemporaryDirectory() as tmp:
            folder=Path(tmp);w=Workbook();s=w.active;s.title='Input Seeds'
            s.append(['ASIN','Marketplace']);s.append(['B000000001','US'])
            p=w.create_sheet('Products');p.append(['ASIN','Marketplace','Product Title','Data Status'])
            p.append(['B000000001','US','Existing product','OK']);w.save(folder/'result.xlsx');w.close()
            (folder/'US_A123456789.json').write_text(json.dumps({'seller':{'seller_id':'A123456789','marketplace':'US','name':'','page_status':'ERROR:TargetClosedError'},'products':[]}),encoding='utf-8')
            state=agent.load_resume(folder)
            self.assertEqual(state['seeds'],[('B000000001','US')])
            self.assertIsNone(state['groups'][0]['seller'])

    def test_amazon_retail_uses_observed_name_without_fabricating_profile(self):
        html=PRODUCT_HTML.replace('Example Seller','Amazon.com')
        with tempfile.TemporaryDirectory() as tmp:
            with patch.object(agent,'sync_playwright'), patch.object(agent,'set_delivery',return_value=('OK','10001')), patch.object(agent,'navigate',return_value=(html,'HTTP_200')) as nav:
                agent.run([('B000000001','US')],tmp,delay=0,skip_storefront=True)
            state=agent.load_resume(tmp);seller=state['groups'][0]['seller']
            self.assertEqual(nav.call_count,1)
            self.assertEqual(seller['name'],'Amazon.com')
            self.assertEqual(seller['status'],'RETAIL_PROFILE_NOT_APPLICABLE')
            self.assertIsNone(seller['visible_rating'])
            self.assertEqual(state['run_status'],'COMPLETED')


if __name__=='__main__':unittest.main()
